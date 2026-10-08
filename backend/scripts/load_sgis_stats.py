"""SGIS 행정구역 통계 zip 을 sgis_area_stats 표에 적재한다 (V071, 세션 453 — 설계서 §5-1).

원천 = 공공데이터포털 15129688 「국가데이터처_SGIS 행정구역 통계 및 경계」 zip(약 269MB, 로그인 없이 내려받기).
사람이 내려받은 zip 경로를 넘긴다(data.go.kr 내려받기는 폼 POST 라 자동화는 후속).

zip 안에서 읽는 것
    * `1. 통계/**.csv` 전부(2024 기준 19개) — 4열 `기준연도,행정구역코드,통계항목,통계값`, cp949.
      행정구역코드는 2글자(시도)·5글자(시군구)·8글자(행정동)가 섞여 있다 → 기본은 8글자만 넣는다
      (`--all-levels` 를 주면 2·5글자도). 값 "N/A"(작은 수 통계 보호)는 value NULL 로 넣는다.
    * `3. 코드집/2. 제공용 코드(statistics_code).xlsx` — 건축년도별 주택 항목(ho_yr_*)의 구간 라벨.
      ho_yr 코드의 뜻은 기준연도마다 다르다(2015년 이후 표: 001 = 1979년 이전 … 020 = 2024년,
      2000·05·10년 표: 001 = 1959년 이전 …). 그래서 ho_yr 행마다 같은 값으로 라벨 항목
      `ho_cy_label_<시작>_<끝>`(예 `ho_cy_label_0000_1979`·`ho_cy_label_2000_2004`·`ho_cy_label_2010_2010`)을
      하나 더 만들고 value_text 에 라벨 원문("1979년 이전")을 담는다 → 화면은 "끝 연도 ≤ 2004" 로 합산.
      코드집에 없는 ho_yr 코드가 나오면 멈춘다(엉뚱한 라벨로 넣지 않는다).

멱등
    (adm_cd, year, item_code) 가 기본키라 `INSERT … ON CONFLICT DO UPDATE` — 두 번 돌려도 행 수·값이 같다
    (loaded_at 만 새 시각). 1,000행 묶음 한 문장씩(운영 연결은 문장당 8초 제한).

사용 (backend 폴더에서)
    python scripts/load_sgis_stats.py <zip경로> --dry-run      # DB 안 씀·연결도 안 함 — 파일별 행 수·항목 수·N/A 수
    python scripts/load_sgis_stats.py <zip경로>                # 2024 행정동(8글자) 적재
    옵션: --year 2024(이 연도 행만, 다른 연도 행은 건너뛰고 센다) · --all-levels(2·5글자도) · --batch-size 1000
⚠ 실제 적재는 V071 이 운영에 적용된 뒤에만. 먼저 --dry-run 이 "멈춤" 없이 끝나는지 본다 —
  적재 중에 멈추면(모르는 건축년도 코드 등) 그 앞 묶음은 이미 들어가 있다(고친 뒤 다시 돌리면 같은 결과).
"""

from __future__ import annotations

import argparse
import csv
import io
import re
import sys
import zipfile
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

HEADER = ["기준연도", "행정구역코드", "통계항목", "통계값"]
STATS_DIR = "/1. 통계/"
CODEBOOK_NAME = "statistics_code"
CODEBOOK_SHEET = "집계구·행정동"
BUILD_YEAR_PREFIX = "ho_yr_"
LABEL_PREFIX = "ho_cy_label_"
NA = "N/A"
DEFAULT_YEAR = 2024
DEFAULT_BATCH = 1000


class SgisLoadError(Exception):
    """적재를 멈춰야 하는 입력 이상(머리줄·코드집·모르는 코드)."""


@dataclass
class FileStats:
    name: str
    rows: int = 0               # 머리줄 뺀 전체 행
    kept: int = 0               # 넣을 행(선택한 단계 + 그 연도)
    by_len: dict[int, int] = field(default_factory=dict)
    na: int = 0                 # 넣을 행 중 N/A
    bad_value: int = 0          # 넣을 행 중 숫자도 N/A 도 아닌 값(value NULL·value_text 원문)
    other_year: int = 0
    items: set[str] = field(default_factory=set)
    label_rows: int = 0         # 만든 건축년도 라벨 행


def parse_value(raw: str) -> tuple[Decimal | None, str | None, str]:
    """통계값 글자 → (value, value_text, 종류). 종류 = ok | na | bad."""
    s = raw.strip()
    if s == NA:
        return None, None, "na"
    try:
        d = Decimal(s)
    except InvalidOperation:
        return None, s, "bad"
    if not d.is_finite():
        return None, s, "bad"
    return d, None, "ok"


def label_item_code(label: str) -> str:
    """건축년도 구간 라벨 → item_code. "1979년 이전" → ho_cy_label_0000_1979,
    "1980년~1989년" → ho_cy_label_1980_1989, "2010년" → ho_cy_label_2010_2010."""
    s = re.sub(r"\s+", "", label)
    m = re.fullmatch(r"(\d{4})년이전", s)
    if m:
        return f"{LABEL_PREFIX}0000_{m.group(1)}"
    m = re.fullmatch(r"(\d{4})년~(\d{4})년", s)
    if m:
        return f"{LABEL_PREFIX}{m.group(1)}_{m.group(2)}"
    m = re.fullmatch(r"(\d{4})년", s)
    if m:
        return f"{LABEL_PREFIX}{m.group(1)}_{m.group(1)}"
    raise SgisLoadError(f"건축년도 라벨을 알아볼 수 없어요: {label!r}")


def load_build_year_labels(xlsx_bytes: bytes, year: int) -> dict[str, str]:
    """코드집 xlsx 에서 그 기준연도에 맞는 건축년도 표(ho_yr 코드 → 라벨)를 고른다.

    시트 '집계구·행정동' 의 열 = (빈칸, 분류, 소분류, 통계항목, 코드, …). 소분류 칸이 차 있으면 새 묶음이다.
    ho_yr 묶음이 둘 있다 — 머리에 '2015' 가 든 표(2015년 이후 기준)와 '2000' 표(2000·05·10년 기준).
    """
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    if CODEBOOK_SHEET not in wb.sheetnames:
        raise SgisLoadError(f"코드집에 '{CODEBOOK_SHEET}' 시트가 없어요: {wb.sheetnames}")
    blocks: list[tuple[str, dict[str, str]]] = []
    header = ""
    for row in wb[CODEBOOK_SHEET].iter_rows(values_only=True):
        cells = list(row) + [None] * 5
        if cells[2] is not None:
            header = str(cells[2])
        code, label = cells[4], cells[3]
        if code is None or not str(code).startswith(BUILD_YEAR_PREFIX):
            continue
        if not blocks or blocks[-1][0] != header:
            blocks.append((header, {}))
        blocks[-1][1][str(code).strip()] = str(label).strip()
    want = "2015" if year >= 2015 else "2000"
    chosen = [m for h, m in blocks if want in h]
    if len(chosen) != 1:
        raise SgisLoadError(
            f"코드집에서 {year}년 기준 건축년도 표를 하나로 고르지 못했어요(찾은 묶음 {[h for h, _ in blocks]})"
        )
    return chosen[0]


def find_members(zf: zipfile.ZipFile) -> tuple[list[str], str]:
    """zip 안 통계 CSV 목록과 코드집 xlsx 이름."""
    csvs = sorted(n for n in zf.namelist() if STATS_DIR in n and n.lower().endswith(".csv"))
    books = [n for n in zf.namelist() if CODEBOOK_NAME in n and n.lower().endswith(".xlsx")]
    if not csvs:
        raise SgisLoadError("zip 안에 '1. 통계/' CSV 가 없어요")
    if len(books) != 1:
        raise SgisLoadError(f"zip 안 코드집(statistics_code xlsx)이 {len(books)}개예요")
    return csvs, books[0]


def iter_file_rows(
    stream: Iterable[bytes] | io.BufferedIOBase,
    stats: FileStats,
    *,
    year: int,
    levels: frozenset[int],
    labels: dict[str, str],
) -> Iterator[dict]:
    """CSV 하나(cp949 바이트 스트림) → 넣을 행 dict. 통계는 stats 에 쌓는다."""
    reader = csv.reader(io.TextIOWrapper(stream, encoding="cp949", newline=""))
    head = next(reader, None)
    if head is None or [h.strip().lstrip("﻿") for h in head] != HEADER:
        raise SgisLoadError(f"{stats.name}: 머리줄이 {HEADER} 가 아니에요 → {head}")
    for line_no, row in enumerate(reader, start=2):
        if not row:
            continue
        if len(row) != 4:
            raise SgisLoadError(f"{stats.name}:{line_no}: 4열이 아니에요 → {row}")
        stats.rows += 1
        y_raw, adm_cd, item, raw = (c.strip() for c in row)
        stats.by_len[len(adm_cd)] = stats.by_len.get(len(adm_cd), 0) + 1
        if len(adm_cd) not in levels or not adm_cd.isdigit():
            continue
        if y_raw != str(year):
            stats.other_year += 1
            continue
        value, value_text, kind = parse_value(raw)
        stats.kept += 1
        stats.items.add(item)
        if kind == "na":
            stats.na += 1
        elif kind == "bad":
            stats.bad_value += 1
        yield {"adm_cd": adm_cd, "year": year, "item_code": item, "value": value, "value_text": value_text}
        if item.startswith(BUILD_YEAR_PREFIX):
            label = labels.get(item)
            if label is None:
                raise SgisLoadError(f"{stats.name}:{line_no}: 코드집에 없는 건축년도 코드 {item}")
            stats.label_rows += 1
            yield {
                "adm_cd": adm_cd, "year": year, "item_code": label_item_code(label),
                "value": value, "value_text": label,
            }


def iter_zip_rows(
    zf: zipfile.ZipFile, *, year: int, levels: frozenset[int], stats_out: list[FileStats]
) -> Iterator[dict]:
    """zip 전체 → 넣을 행. 파일마다 FileStats 를 stats_out 에 붙인다."""
    csvs, book = find_members(zf)
    labels = load_build_year_labels(zf.read(book), year)
    for name in csvs:
        st = FileStats(name=name.rsplit("/", 1)[-1])
        stats_out.append(st)
        with zf.open(name) as f:
            yield from iter_file_rows(f, st, year=year, levels=levels, labels=labels)


def _batched(rows: Iterable[dict], size: int) -> Iterator[list[dict]]:
    """size 행씩 묶되, 한 묶음 안에서 같은 키는 마지막 값 하나만 남긴다
    (PostgreSQL 은 한 문장에서 같은 행을 두 번 고치면 거부한다)."""
    batch: dict[tuple, dict] = {}
    for r in rows:
        batch[(r["adm_cd"], r["year"], r["item_code"])] = r
        if len(batch) >= size:
            yield list(batch.values())
            batch = {}
    if batch:
        yield list(batch.values())


def upsert_rows(db, rows: Iterable[dict], batch_size: int = DEFAULT_BATCH) -> int:
    """INSERT … ON CONFLICT (adm_cd, year, item_code) DO UPDATE — PostgreSQL·SQLite 둘 다.

    묶음마다 커밋한다(중간에 끊겨도 다시 돌리면 이어서 같은 결과). 넣은(고친) 행 수를 돌려준다.
    """
    from db.models import SgisAreaStats
    from utils import utcnow

    if db.get_bind().dialect.name == "sqlite":
        from sqlalchemy.dialects.sqlite import insert as db_insert
    else:
        from sqlalchemy.dialects.postgresql import insert as db_insert

    total = 0
    for batch in _batched(rows, batch_size):
        now = utcnow()
        for r in batch:
            r["loaded_at"] = now
        stmt = db_insert(SgisAreaStats).values(batch)
        stmt = stmt.on_conflict_do_update(
            index_elements=["adm_cd", "year", "item_code"],
            set_={
                "value": stmt.excluded.value,
                "value_text": stmt.excluded.value_text,
                "loaded_at": stmt.excluded.loaded_at,
            },
        )
        db.execute(stmt)
        db.commit()
        total += len(batch)
    return total


def _print_stats(stats: list[FileStats]) -> None:
    for s in stats:
        lens = " ".join(f"{k}글자 {v:,}" for k, v in sorted(s.by_len.items()))
        print(
            f"- {s.name}: 전체 {s.rows:,}행({lens}) → 넣을 행 {s.kept:,} · 항목 {len(s.items)}개"
            f" · N/A {s.na:,} · 숫자 아님 {s.bad_value:,} · 다른 연도 {s.other_year:,} · 건축년도 라벨 {s.label_rows:,}"
        )
    kept = sum(s.kept for s in stats)
    labels = sum(s.label_rows for s in stats)
    print(f"합계: 파일 {len(stats)}개 · 넣을 행 {kept:,} + 라벨 행 {labels:,} = {kept + labels:,}"
          f" · N/A {sum(s.na for s in stats):,} · 숫자 아님 {sum(s.bad_value for s in stats):,}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="SGIS 행정구역 통계 zip → sgis_area_stats 적재")
    parser.add_argument("zip_path", help="공공데이터포털 15129688 zip 경로")
    parser.add_argument("--year", type=int, default=DEFAULT_YEAR, help=f"넣을 기준연도 (기본 {DEFAULT_YEAR})")
    parser.add_argument("--all-levels", action="store_true", help="시도(2글자)·시군구(5글자) 행도 넣기 (기본은 행정동 8글자만)")
    parser.add_argument("--dry-run", action="store_true", help="DB 를 쓰지도 연결하지도 않고 파일별 행 수만 센다")
    parser.add_argument("--batch-size", type=int, default=DEFAULT_BATCH, help=f"한 문장에 넣을 행 수 (기본 {DEFAULT_BATCH})")
    args = parser.parse_args(argv)

    levels = frozenset({2, 5, 8}) if args.all_levels else frozenset({8})
    stats: list[FileStats] = []
    try:
        with zipfile.ZipFile(args.zip_path) as zf:
            rows = iter_zip_rows(zf, year=args.year, levels=levels, stats_out=stats)
            if args.dry_run:
                for _ in rows:
                    pass
                written = 0
            else:
                from db.database import SessionLocal

                with SessionLocal() as db:
                    written = upsert_rows(db, rows, args.batch_size)
    except SgisLoadError as e:
        print(f"멈춤: {e}")
        return 2

    _print_stats(stats)
    if sum(s.kept for s in stats) == 0:
        print(f"멈춤: {args.year}년 행이 하나도 없어요 — --year 를 확인하세요")
        return 2
    if args.dry_run:
        print("(시험 실행 — DB 에 쓰지 않았어요)")
    else:
        print(f"적재 끝: {written:,}행 넣음(이미 있던 행은 값만 고침)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
